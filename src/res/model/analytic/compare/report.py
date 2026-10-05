"""Compact matplotlib reports and typed Excel tables using project dependencies."""
from __future__ import annotations

import json
import re
import textwrap
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from src.proj.util.functional.plot import (
    new_figure, plot_table, set_seaborn_theme, set_xaxis, set_yaxis,
)


def make_highlights(result):
    messages = []
    for metric, title in [('ic_mean', 'Highest mean IC'), ('top_excess_sum', 'Highest cumulative daily excess'),
                          ('top_excess_annualized', 'Highest annualized excess return')]:
        if metric in result.summary:
            values = result.summary.set_index('model')[metric].dropna()
            if len(values):
                messages.append(f'{title}: {values.idxmax()} ({values.max():.4f}).')
    for key, title in [('corr_ic_all', 'IC'), ('corr_excess_all', 'daily excess return')]:
        matrix = result.correlations.get(key)
        if matrix is None:
            continue
        pairs = [(matrix.iloc[i, j], matrix.index[i], matrix.columns[j])
                 for i in range(len(matrix)) for j in range(i + 1, len(matrix)) if pd.notna(matrix.iloc[i, j])]
        if pairs:
            low, high = min(pairs), max(pairs)
            messages.append(f'{title} correlation: lowest {low[1]} / {low[2]} ({low[0]:.3f}); '
                            f'highest {high[1]} / {high[2]} ({high[0]:.3f}).')
    for kind, metric in [('ic', 'ic_mean'), ('top', 'excess_annualized')]:
        frame = result.tables.get(f'{kind}_periods')
        if frame is None or frame.empty:
            continue
        full = frame[frame.period == 'all'].set_index('model')[metric]
        recent = frame[frame.period == 'recent_year'].set_index('model')[metric]
        for model, delta in (recent - full).dropna().items():
            messages.append(f'{model}: recent-year {metric} versus full window {delta:+.4f}.')
    return messages


def _figure(rows: int = 1, height: float = 7):
    fig = new_figure(size=(16, height))
    axes = fig.subplots(rows, 1, squeeze=False)[:, 0]
    return fig, axes


def _heatmap(matrix, title, cka=False):
    import seaborn as sns
    fig, (ax,) = _figure(height=max(7, len(matrix) * .55))
    values = matrix.to_numpy(dtype=float)
    cmap = sns.light_palette('seagreen', as_cmap=True) if cka else sns.diverging_palette(140, 10, sep=10, as_cmap=True)
    sns.heatmap(matrix, ax=ax, vmin=0 if cka else -1, vmax=1, cmap=cmap,
                annot=len(matrix) <= 12, fmt='.2f', annot_kws={'fontsize': 10},
                square=True, linewidths=.5, cbar_kws={'shrink': .8})
    labels = [textwrap.fill(str(x), 26) for x in matrix.index]
    set_xaxis(ax, np.arange(len(matrix)) + .5, labels=labels, title='Model', grid=False)
    ax.set_yticks(np.arange(len(matrix)) + .5, labels, rotation=0)
    set_yaxis(ax, format='default', title='Model')
    ax.set_title(title, fontsize=14)
    if len(matrix) <= 12:
        for i in range(len(matrix)):
            for j in range(len(matrix)):
                if np.isnan(values[i, j]):
                    ax.text(j + .5, i + .5, 'NA', ha='center', va='center', fontsize=10)
    return fig


def build_figures(result):
    import matplotlib as mpl
    import matplotlib.pyplot as plt
    import seaborn as sns
    set_seaborn_theme()
    names = result.summary.model.tolist()
    colors = sns.diverging_palette(140, 10, sep=10, n=len(names))
    # The project's diverging palette has a near-white midpoint for odd counts.
    colors = [tuple(channel * .6 for channel in color) if min(color) > .8 else color for color in colors]
    palette = dict(zip(names, colors))
    # Use the same theme, table renderer and axis helpers as factor/top plotters.
    with mpl.rc_context():
        summary = result.summary.copy()
        # Split the summary into readable pages for many models.
        columns = [c for c in ('model', 'ic_mean', 'top_pf_return', 'top_excess_sum', 'top_excess_annualized', 'top_ir', 'top_excess_mdd') if c in summary]
        for offset in range(0, len(summary), 12):
            fig, (ax, notes_ax) = _figure(rows=2)
            ax.set_title('Model Comparison Front Face', fontsize=14)
            frame = summary.iloc[offset:offset + 12][columns].copy()
            for col in frame:
                percent = col in ('top_pf_return', 'top_excess_sum', 'top_excess_annualized', 'top_excess_mdd')
                frame[col] = frame[col].map(lambda x: 'NA' if pd.isna(x) else
                                           (f'{x:.2%}' if percent else f'{x:.4f}') if isinstance(x, float) else str(x))
            frame['model'] = frame.model.map(lambda x: textwrap.fill(x, 30))
            labels = {'model': 'Model', 'ic_mean': 'Mean IC', 'ic_n': 'IC N',
                      'top_pf_return': 'PF return', 'top_excess_sum': 'Excess sum',
                      'top_excess_annualized': 'Excess ann.', 'top_ir': 'IR', 'top_excess_mdd': 'Excess MDD'}
            plt.sca(ax)
            plot_table(frame.rename(columns=labels).set_index('Model'), capitalize=False,
                       fontsize=10, index_width=3.5, stripe_by=1, column_definitions=[])
            notes = [f'{kind.removesuffix("_window").upper()}: {result.metadata[kind]["start"]} - '
                     f'{result.metadata[kind]["end"]} ({result.metadata[kind]["n_dates"]} dates)'
                     for kind in ('ic_window', 'top_window') if kind in result.metadata]
            notes += ['Daily excess is additive; portfolio/benchmark returns compound.',
                      'Full statistics, observation counts and diagnostics are in Excel.',
                      f'Diagnostics: {len(result.diagnostics)}; optional sample records: {len(result.samples)}.']
            if not result.diagnostics.empty:
                notes.extend(result.diagnostics.reason.drop_duplicates().tolist()[:3])
            notes_ax.axis('off')
            notes_ax.text(0, 1, '\n'.join(textwrap.fill(n, 120) for n in notes), va='top',
                          transform=notes_ax.transAxes, fontsize=10)
            result.figures[f'summary_{offset // 12 + 1}'] = fig
        for kind in ('ic', 'top'):
            frame = result.tables.get(f'{kind}_series')
            if frame is None:
                continue
            specs = [('ic', 'Raw IC'), ('rolling_ic', f'Rolling IC ({result.metadata["config"]["rolling_window"]} observations)'),
                     ('cum_ic', 'Cumulative IC')] if kind == 'ic' else [
                         ('cum_pf', 'Portfolio compound return'), ('cum_excess', 'Cumulative daily excess'),
                         ('excess_drawdown', 'Excess drawdown')]
            fig, axes = _figure(rows=3, height=12)
            for ax, (column, title) in zip(axes, specs):
                for model, group in frame.groupby('model', sort=False):
                    times = group.date.astype(str)
                    plotted = group[column].copy()
                    if kind == 'top':
                        plotted.loc[group.path_gap] = np.nan
                    ax.plot(times, plotted, label=model, color=palette[model])
                    if kind == 'top':
                        endpoints = group.path_gap & np.isfinite(group[column])
                        if endpoints.any():
                            ax.scatter(times[endpoints], group.loc[endpoints, column], s=12, color=palette[model])
                ax.set_title(f'Model Comparison {title}', fontsize=14)
                set_xaxis(ax, frame.date.astype(str).unique(), title='Trade Date')
                set_yaxis(ax, format='pct' if kind == 'top' else 'flt', digits=2,
                          title=title, title_color='b')
                ax.legend(loc='upper left', fontsize=10, ncol=min(3, len(names)))
            result.figures['ic_trend' if kind == 'ic' else 'top_curves'] = fig
            frame = result.tables[f'{kind}_periods']
            frame = frame[frame.period != 'all']
            metric = 'ic_mean' if kind == 'ic' else 'excess_annualized'
            # Paginate period bars instead of squeezing every month onto one page.
            labels = frame.period.unique()
            for offset in range(0, len(labels), 12):
                subset = frame[frame.period.isin(labels[offset:offset + 12])]
                pivot = subset.pivot(index='period', columns='model', values=metric).reindex(labels[offset:offset + 12])
                pivot = pivot.reindex(columns=names)
                fig, (ax,) = _figure()
                pivot.plot.bar(ax=ax, width=.8, color=[palette[name] for name in pivot.columns])
                set_xaxis(ax, np.arange(len(pivot)), labels=pivot.index, title='Period')
                set_yaxis(ax, format='pct' if kind == 'top' else 'flt', digits=2 if kind == 'top' else 3,
                          title='Annualized Excess Return' if kind == 'top' else 'Average IC', title_color='b')
                for position, (_, values) in enumerate(pivot.iterrows()):
                    if values.isna().all():
                        ax.text(position, .025, 'NA', ha='center', transform=ax.get_xaxis_transform())
                    elif values.isna().any():
                        for model_index, missing in enumerate(values.isna()):
                            if missing:
                                x = position - .4 + .8 * (model_index + .5) / len(values)
                                ax.text(x, .025, 'NA', ha='center', rotation=90, fontsize=6,
                                        transform=ax.get_xaxis_transform())
                ax.set_title(f'Model Comparison {kind.upper()} by Period', fontsize=14)
                ax.legend(loc='upper left', fontsize=10, ncol=min(3, len(names)))
                result.figures[f'{kind}_periods_{offset // 12}'] = fig
        for key, matrix in result.correlations.items():
            if key.endswith('_count'):
                continue
            # PDF focuses on main matrices; all period matrices are in Excel.
            if key in ('corr_ic_all', 'corr_excess_all', 'corr_pf_all', 'pred_pearson_mean',
                       'pred_spearman_mean', 'hidden_cka_mean'):
                result.figures[key] = _heatmap(matrix, key.replace('_', ' ').title(), key == 'hidden_cka_mean')
        for fig in result.figures.values():
            fig.tight_layout()
            plt.close(fig)


def export_report(result, output_dir=None):
    from src.proj import PATH, Save
    if output_dir is None:
        output_dir = PATH.result / 'model_compare' / datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    directory = Path(output_dir).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    excel, pdf = directory / 'comparison.xlsx', directory / 'comparison.pdf'
    correlation_tables = {}
    for prefix in ('corr_ic_', 'corr_excess_', 'corr_pf_', 'pred_pearson_', 'pred_spearman_', 'hidden_cka_'):
        blocks = []
        for key, matrix in result.correlations.items():
            if not key.startswith(prefix) or key.endswith('_count'):
                continue
            count_key = f'{key}_count' if key.startswith('corr_') else key.removesuffix('mean') + 'count'
            counts = result.correlations[count_key].add_prefix('N:')
            block = matrix.join(counts).rename_axis('model').reset_index()
            block.insert(0, 'period', key.removeprefix(prefix))
            blocks.append(block)
        if blocks:
            correlation_tables[prefix.rstrip('_')] = pd.concat(blocks, ignore_index=True)
    detail_tables = {key: value for key, value in result.tables.items() if key != 'hidden_dimensions'}
    hidden = result.tables.get('hidden_dimensions')
    if hidden is not None and not hidden.empty:
        for date, group in hidden.groupby('date', sort=True):
            detail_tables[f'hidden_{date}'] = group.sort_values(['model_a', 'model_b', 'feature_a', 'feature_b'])
    tables = {'Summary': result.summary, 'Highlights': pd.DataFrame({'finding': result.highlights}), **detail_tables,
              **correlation_tables,
              'Samples': result.samples, 'Diagnostics': result.diagnostics,
              'Sources': pd.DataFrame(result.metadata['sources']),
              'Methodology': pd.DataFrame([{'key': k, 'value': json.dumps(v, ensure_ascii=False, default=str)}
                                          for k, v in result.metadata.items() if k != 'sources'])}
    names, chunks, mapping = set(), {}, []
    for key, frame in tables.items():
        if frame.empty:
            continue
        # Excel supports 1,048,576 rows including the header; preserve every row.
        for offset in range(0, len(frame), 1_000_000):
            base = re.sub(r'[\[\]:*?/\\]', '_', key)[:25]
            name, number = base, 1
            while name in names:
                number += 1
                name = f'{base}_{number}'
            names.add(name)
            chunks[name] = frame.iloc[offset:offset + 1_000_000].reset_index(drop=True)
            mapping.append({'table': key, 'sheet': name, 'row_offset': offset})
    chunks['Sheet_index'] = pd.DataFrame(mapping)
    # Synchronous project exports: errors propagate and completion means files exist.
    Save.dfs(chunks, excel, async_save=False)
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties
    book = load_workbook(excel)
    for sheet in book:
        sheet.freeze_panes = 'C2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='233B53')
        for col_num, cells in enumerate(sheet.iter_cols(), 1):
            preview = list(cells[:min(len(cells), 100)])
            width = min(52, max(12, max(len(str(c.value or '')) for c in preview) + 2))
            sheet.column_dimensions[get_column_letter(col_num)].width = width
        for row in sheet.iter_rows(min_row=2):
            for column, cell in enumerate(row, 1):
                if isinstance(cell.value, float):
                    header = str(sheet.cell(1, column).value)
                    percent = any(header.endswith(x) for x in ('pf_return', 'bm_return', 'return_difference',
                                  'excess_sum', 'excess_annualized', 'tracking_error', 'excess_mdd', 'positive_rate'))
                    cell.number_format = '0.00%' if percent else '0.0000'
                elif isinstance(cell.value, str) and len(cell.value) > 52:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    book.save(excel)
    Save.figs(result.figures, pdf, async_save=False, close=False)
    parameters = directory / 'compare_parameters.json'
    parameters.write_text(json.dumps({
        'exported_at': datetime.now().astimezone().isoformat(),
        'output_dir': str(directory),
        **result.metadata,
    }, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
    result.output_paths = {'xlsx': excel, 'pdf': pdf, 'parameters': parameters}
